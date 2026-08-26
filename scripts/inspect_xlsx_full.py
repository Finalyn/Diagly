"""Full inspection: sample distribution + structural patterns of both sheets."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from openpyxl import load_workbook
from collections import Counter

path = r"C:\Users\switc\Desktop\DIAGLY_FINAL.xlsx"
wb = load_workbook(path, data_only=True, read_only=True)

# ---------- Feuil1 ----------
ws1 = wb["Feuil1"]
rows1 = list(ws1.iter_rows(values_only=True))
print("=" * 80)
print(f"FEUIL1 — {len(rows1)} rows")
print("=" * 80)

# Show rows 0..2 verbatim (header zone)
for i, r in enumerate(rows1[:3]):
    print(f"  row{i}: {r}")
print()

# Categorize by code shape
def classify(code):
    if code is None:
        return "null"
    s = str(code).strip()
    if not s:
        return "empty"
    if s.replace(".", "").isdigit():
        if "." not in s:
            n = len(s)
            return f"int-{n}d"  # 1, 101, etc.
        else:
            parts = s.split(".")
            return f"dec-{len(parts[0])}.{len(parts[1])}"
    return f"other:{s[:20]}"

shapes = Counter()
non_empty_state_rows = 0
for r in rows1[2:]:  # skip both headers
    code = r[0]
    shapes[classify(code)] += 1
    # any state column filled?
    if any(c is not None and str(c).strip() for c in r[2:]):
        non_empty_state_rows += 1

print("Code shape distribution:")
for k, v in shapes.most_common():
    print(f"  {k:>15s}: {v}")
print(f"\nRows with at least one state column filled: {non_empty_state_rows}")

# Show some non-empty-state rows from Feuil1
print("\nFEUIL1 — rows that DO have state descriptions (first 15):")
shown = 0
for i, r in enumerate(rows1[2:], start=2):
    if any(c is not None and str(c).strip() for c in r[2:]):
        print(f"  row{i}: code={r[0]!r} label={r[1]!r}")
        for j, label in enumerate(['TBE', 'Bon', 'Moyen', 'Mauvais', 'Amelio', 'Normes'], start=2):
            v = r[j]
            if v is not None and str(v).strip():
                print(f"    {label}: {v!r}")
        shown += 1
        if shown >= 15:
            break

# ---------- Feuil2 ----------
print("\n" + "=" * 80)
print(f"FEUIL2 — {len(list(wb['Feuil2'].iter_rows()))} rows")
print("=" * 80)

ws2 = wb["Feuil2"]
rows2 = list(ws2.iter_rows(values_only=True))

# Print ALL rows of Feuil2 since only 142
print("\nFEUIL2 full dump (header + all 142 rows):\n")
header = rows2[0]
for i, r in enumerate(rows2):
    print(f"row{i}:")
    for j, v in enumerate(r):
        if v is not None and str(v).strip():
            col = header[j] if i > 0 and j < len(header) else f"col{j}"
            print(f"  [{j}] {col!r}: {v!r}")
    print()
