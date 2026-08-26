"""Inspect the structure of DIAGLY_FINAL.xlsx — sheets, columns, sample rows."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\switc\Desktop\DIAGLY_FINAL.xlsx"
wb = load_workbook(path, data_only=True, read_only=True)

print(f"FILE: {path}")
print(f"SHEETS ({len(wb.sheetnames)}): {wb.sheetnames}\n")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print("=" * 80)
    print(f"SHEET: {sheet_name!r}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print("=" * 80)

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("  (empty)")
        continue

    header = rows[0]
    print(f"HEADER ({len(header)} cols):")
    for i, h in enumerate(header):
        print(f"  [{i}] {h!r}")

    n_sample = min(5, len(rows) - 1)
    if n_sample > 0:
        print(f"\nFIRST {n_sample} DATA ROWS:")
        for r in rows[1:1 + n_sample]:
            print(f"  {r}")

    print(f"\nTOTAL DATA ROWS: {len(rows) - 1}\n")
